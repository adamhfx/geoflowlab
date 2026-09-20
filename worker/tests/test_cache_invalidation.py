import zipfile
import xml.etree.ElementTree as ET
import openpyxl
from worker.engine import inject_inputs

def test_input_injection_removes_all_formula_caches(tmp_path):
    original=tmp_path/'original.xlsx';seed=tmp_path/'cached.xlsx';out=tmp_path/'out.xlsx'
    wb=openpyxl.Workbook();wb.active.title='Inputs';wb.active['A1']=10
    wb.create_sheet('Results')['A1']='Inputs!A1*2'
    wb['Results']['B1']='=Inputs!A1*3';wb.save(original)
    ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(original) as src,zipfile.ZipFile(seed,'w') as dst:
        for name in src.namelist():
            data=src.read(name)
            if name=='xl/worksheets/sheet2.xml':
                root=ET.fromstring(data)
                root.find('.//m:c[@r="B1"]/m:v',ns).text='30'
                data=ET.tostring(root)
            dst.writestr(name,data)
    inject_inputs(seed,out,{'A1':20})
    with zipfile.ZipFile(out) as src:
        root=ET.fromstring(src.read('xl/worksheets/sheet2.xml'))
        assert root.find('.//m:c[@r="B1"]/m:f',ns).text=='Inputs!A1*3'
        assert root.find('.//m:c[@r="B1"]/m:v',ns) is None
    assert openpyxl.load_workbook(out,data_only=True)['Inputs']['A1'].value==20
