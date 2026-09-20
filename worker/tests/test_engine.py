import json
from pathlib import Path

import openpyxl
import pytest

from worker.engine import ModelError, extract, inject_inputs, payout_year, validate_inputs


def manifest():
    return {"fields": [
        {"id": "rate", "label": "Rate", "type": "number", "required": True, "min": 0, "exclusiveMin": 0},
        {"id": "note", "label": "Note", "type": "text", "required": False},
        {"id": "m1", "label": "Completion 1", "type": "milestone", "required": False},
        {"id": "m2", "label": "Completion 2", "type": "milestone", "required": False},
    ]}


@pytest.mark.parametrize(("flows", "expected"), [
    ([-10, -2, -1], None),                 # never recovers
    ([5, -8, 1, 2], 4),                    # initially positive, then deficit
    ([-3, 3], 2),                          # recovery boundary
    ([0, 0, 2], 0),                        # no deficit means no payout period
])
def test_payout_year_edges(flows, expected):
    assert payout_year(flows) == expected


def test_validate_required_zero_and_injection_inputs():
    m = manifest()
    with pytest.raises(ModelError):
        validate_inputs({"rate": 0, "m1": "X", "m2": ""}, m)
    with pytest.raises(ModelError):
        validate_inputs({"rate": 1, "m1": "X", "m2": "", "unknown": 2}, m)
    result = validate_inputs({"rate": 1, "note": "=HYPERLINK(\"https://evil\")", "m1": "X", "m2": ""}, m)
    assert result["note"].startswith("=")


def test_inject_inputs_never_overwrites_formula(tmp_path):
    source = tmp_path / "template.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws["A1"] = "safe"
    ws["B1"] = "=1+1"
    wb.save(source)
    with pytest.raises(ModelError, match="formula"):
        inject_inputs(source, tmp_path / "out.xlsx", {"A1": "=literal", "B1": 4})
    inject_inputs(source, tmp_path / "out.xlsx", {"A1": "=literal"})
    check = openpyxl.load_workbook(tmp_path / "out.xlsx", data_only=False)
    assert check["Inputs"]["A1"].value == "=literal"
    assert check["Inputs"]["A1"].data_type == "s"


def test_extract_normalizes_blank_late_capex_only():
    class Cell:
        def __init__(self, value=1):
            self.value, self.data_type = value, 'n'
    class Sheet:
        def __init__(self, cash=False):
            self.cash = cash; self.cells = {}
        def __getitem__(self, address):
            if address not in self.cells:
                col=''.join(x for x in address if x.isalpha()); row=int(''.join(x for x in address if x.isdigit()))
                value=1
                if self.cash and col == 'W' and row >= 35: value=None
                if self.cash and address == 'AL44': value=10
                if self.cash and address == 'AN44': value=20
                if self.cash and address == 'AL41': value=5
                if self.cash and address == 'W41': value=100
                if self.cash and address == 'AL52': value=.1
                self.cells[address]=Cell(value)
            return self.cells[address]
        def __setitem__(self, address, value): self.cells[address] = value if isinstance(value, Cell) else Cell(value)
        def __iter__(self): return iter(())
    class Book:
        def __init__(self): self.sheets={'CashFlowSummStRnt':Sheet(True),'DashboardStRnt':Sheet()}
        def __getitem__(self, name): return self.sheets[name]
    result=extract(Book())
    assert result['annual'][-1]['capex'] == 0
