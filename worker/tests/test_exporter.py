import zipfile
import xml.etree.ElementTree as ET

import openpyxl

from worker.exporter import inspect_export, export_results


def test_export_is_values_only_and_blocks_formula_strings(tmp_path):
    source = openpyxl.Workbook()
    source.active.title = "CashFlowSummStRnt"
    source.create_sheet("DashboardStRnt")
    for sheet in source.worksheets:
        sheet["A1"] = 0
        sheet["B1"] = "=HYPERLINK(\"https://example.invalid\")"
        # Model an untrusted text value that happens to begin with '='.
        sheet["B1"].data_type = "s"
        sheet.merge_cells("C1:D1")
        sheet["C1"] = 0
        for row in range(7, 37):
            for col in range(1, 46):
                sheet.cell(row, col).value = row + col
    output = tmp_path / "export.xlsx"
    export_results(source, {"sensitivity": [{"label": "Case", "low": 1, "high": 2, "baseline": 1}]}, output)
    assert inspect_export(output)
    with zipfile.ZipFile(output) as archive:
        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        for name in archive.namelist():
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
                root = ET.fromstring(archive.read(name))
                assert not root.findall(".//s:f", ns)
    check = openpyxl.load_workbook(output, data_only=False)
    assert check["CashFlowSummStRnt"]["B1"].value == "=HYPERLINK(\"https://example.invalid\")"
    assert check["CashFlowSummStRnt"]["C1"].value == 0
